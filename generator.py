"""
Builds the Programme of Research Work (Acad Form No. 19, KVASU) as a .docx
following the 'General instructions while preparing the Programme of
Research work' and the modified form ordered on 30.10.2023.

Layout rules implemented
  * A4, 2.5 cm margins all sides, Times New Roman 12, line spacing 1.15
  * Top headings in one column; everything else in two columns of 7.75 cm
  * One blank line before every major serial number (1-14)
  * Paragraphs start with a 1 cm tab where required, matter justified
  * No colon after 'Name of the discipline' / 'Project coordination group'
  * References justified with hanging indent
  * 'Sd/-' added for soft copy (optional)

Inline formatting inside any text field:
  <i>italic</i>  <b>bold</b>  <sup>superscript</sup>  <sub>subscript</sub>
"""
from __future__ import annotations

import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

try:
    from .references import format_reference, sort_key
except ImportError:  # files placed flat in one folder
    from references import format_reference, sort_key

FONT = "Times New Roman"
COL_W = 7.75  # cm
HANG = 0.6    # cm, indent for numbered headings / lists
J, L, R, C = (WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT,
              WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER)

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
LETTERS = "abcdefghijklmnopqrstuvwxyz"
SMALL_ROMAN = ["i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x",
               "xi", "xii", "xiii", "xiv", "xv"]


# ---------------------------------------------------------------- helpers
_TAG = re.compile(r"(</?(?:i|b|sup|sub)>)")


def add_rich(par, text: str, bold=False, italic=False):
    """Add text with <i>, <b>, <sup>, <sub> markup to a paragraph."""
    state = {"i": italic, "b": bold, "sup": False, "sub": False}
    for tok in _TAG.split(text or ""):
        m = re.fullmatch(r"<(/?)(i|b|sup|sub)>", tok)
        if m:
            state[m.group(2)] = not m.group(1)
            continue
        if not tok:
            continue
        run = par.add_run(tok)
        _font(run, bold=state["b"], italic=state["i"])
        run.font.superscript = state["sup"] or None
        run.font.subscript = state["sub"] or None
    return par


def _font(run, bold=False, italic=False, size=12):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)


class Builder:
    def __init__(self, data: dict):
        self.d = data
        self.sd = data.get("include_sd", True)
        self.doc = Document()
        self._setup()

    # ------------------------------------------------------------ setup
    def _setup(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(12)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        pf = st.paragraph_format
        pf.line_spacing = 1.15
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.widow_control = True
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, side, Cm(2.5))

    def p(self, text="", align=J, bold=False, italic=False, first=None,
          left=None, hang=None, keep_next=False):
        par = self.doc.add_paragraph()
        par.alignment = align
        f = par.paragraph_format
        if left is not None:
            f.left_indent = Cm(left)
        if first is not None:
            f.first_line_indent = Cm(first)
        if hang is not None:
            f.left_indent = Cm((left or 0) + hang)
            f.first_line_indent = Cm(-hang)
        if keep_next:
            f.keep_with_next = True
        if text:
            add_rich(par, text, bold=bold, italic=italic)
        return par

    def blank(self):
        return self.p("", align=L)

    def heading(self, text, hang=HANG, left=0, align=L):
        """Bold sub-heading such as '5. Objectives of the study'."""
        return self.p(text, align=align, bold=True, hang=hang, left=left, keep_next=True)

    def right_tab(self, par, pos_cm=COL_W):
        par.paragraph_format.tab_stops.add_tab_stop(Cm(pos_cm), WD_TAB_ALIGNMENT.RIGHT)

    # ------------------------------------------------------------ build
    def build(self, out_path: str):
        d = self.d
        degree = d.get("degree", "MASTERS").upper()
        # ---- top headings (single column) ----
        self.p("Acad. Form No. 19", align=R, bold=True)
        self.blank()
        self.p("KERALA VETERINARY AND ANIMAL SCIENCES UNIVERSITY", align=C, bold=True)
        self.p(d.get("faculty", "Faculty of Veterinary and Animal Sciences"), align=C, bold=True)
        top = self.p(f"PROGRAMME OF RESEARCH WORK FOR THESIS FOR {degree} DEGREE",
                     align=C, bold=True)
        # keep it on one line: tighten letter spacing slightly if long
        if len(top.text) > 56:
            for r in top.runs:
                sp = OxmlElement("w:spacing")
                sp.set(qn("w:val"), "-6")
                r._element.get_or_add_rPr().append(sp)

        # ---- two-column body ----
        sec = self.doc.add_section(WD_SECTION.CONTINUOUS)
        cols = sec._sectPr.xpath("./w:cols")
        col = cols[0] if cols else OxmlElement("w:cols")
        col.set(qn("w:num"), "2")
        col.set(qn("w:space"), str(int(Cm(21 - 5 - 2 * COL_W).twips)))
        col.set(qn("w:equalWidth"), "1")
        if not cols:
            sec._sectPr.append(col)

        self.blank()
        self.s1_title()
        self.blank(); self.s2_project()
        self.blank(); self.s3_student()
        self.blank(); self.s4_advisor()
        self.blank(); self.s5_objectives()
        self.blank(); self.s6_utility()
        self.blank(); self.s7_publications()
        self.blank(); self.s8_technical()
        self.blank(); self.s9_observations()
        self.blank(); self.s10_facilities()
        self.blank(); self.s11_duration()
        self.blank(); self.s12_finance()
        self.signatures()
        self.blank(); self.committee()
        self.blank(); self.s13_references()
        self.blank(); self.s14_timeframe()
        self.blank(); self.certificate()
        self.doc.save(out_path)
        return out_path

    # ------------------------------------------------------------ sections
    def s1_title(self):
        self.heading("1. Title of thesis")
        self.p(self.d.get("title", ""), align=J, left=HANG)

    def s2_project(self):
        d = self.d
        self.p("2. a. Title of the departmental/KVASU research project of which "
               "this forms a part", align=J, bold=True, hang=HANG, keep_next=True)
        self.p(d.get("project_title", "Nil") or "Nil", align=L, left=HANG)
        self.p("b. Code No. if any, and order by which departmental/KVASU "
               "research project is approved", align=J, bold=True, left=HANG,
               hang=HANG, keep_next=True)
        self.p(d.get("project_code", "Not applicable") or "Not applicable",
               align=L, left=HANG)

    def s3_student(self):
        d = self.d
        self.heading("3. a. Name of student")
        self.p(d.get("student_name", ""), align=L, left=HANG)
        self.heading("b. Admission No.", left=HANG)
        self.p(d.get("admission_no", ""), align=L, left=HANG)
        self.heading("c. Name of the discipline", left=HANG)   # no colon (order 30.10.2023)
        self.p(d.get("discipline", ""), align=L, left=HANG)

    def s4_advisor(self):
        d = self.d
        self.heading("4. a. Name of major advisor")
        self.p(d.get("advisor_name", ""), align=L, left=HANG)
        self.heading("b. Designation", left=HANG)
        for line in _lines(d.get("advisor_designation", "")):
            self.p(line, align=L, left=HANG)

    def s5_objectives(self):
        self.heading("5. Objectives of the study")
        for i, obj in enumerate(self.d.get("objectives", []), 1):
            self.p(f"{i}. {obj.strip().rstrip('.')}", align=J, hang=HANG)

    def s6_utility(self):
        self.heading("6. Practical / Scientific utility")
        for para in _paras(self.d.get("utility", "")):
            self.p(para, align=J, first=1.0)

    def s7_publications(self):
        self.heading("7. Important publications on which the study is based")
        paras = _paras(self.d.get("publications", ""))
        for i, para in enumerate(paras):
            self.p(para, align=J, first=1.0)
            if i < len(paras) - 1:
                pass

    def s8_technical(self):
        self.heading("8. Outline of the technical programme")
        for block in self.d.get("technical", []):
            if isinstance(block, str):
                block = {"type": "para", "text": block}
            if block.get("type") == "table":
                self._table(block)
            else:
                for para in _paras(block.get("text", "")):
                    self.p(para, align=J, first=1.0)

    def _table(self, block):
        headers, rows = block.get("headers", []), block.get("rows", [])
        if block.get("caption"):
            self.p(block["caption"], align=C, bold=True, keep_next=True)
        n = len(headers)
        t = self.doc.add_table(rows=1 + len(rows), cols=n)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        total = COL_W - 0.25
        w = Cm(total / n)
        tblPr = t._tbl.tblPr
        for tag, attrs in (("w:tblW", {"w:w": str(int(Cm(total).twips)), "w:type": "dxa"}),
                           ("w:tblLayout", {"w:type": "fixed"})):
            el = tblPr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag); tblPr.append(el)
            for k, v in attrs.items():
                el.set(qn(k), v)
        for gc in t._tbl.tblGrid.findall(qn("w:gridCol")):
            gc.set(qn("w:w"), str(int(w.twips)))
        for ri, row in enumerate([headers] + rows):
            for ci in range(n):
                cell = t.cell(ri, ci)
                cell.width = w
                cell.text = ""
                par = cell.paragraphs[0]
                par.alignment = C
                par.paragraph_format.line_spacing = 1.0
                add_rich(par, str(row[ci]) if ci < len(row) else "", bold=(ri == 0))
        self.blank()

    def s9_observations(self):
        self.heading("9. Main items of observations to be made")
        for i, grp in enumerate(self.d.get("observations", []), 1):
            grp = _norm_item(grp)
            self.p(f"{i}. {grp['text']}", align=L, hang=HANG, keep_next=bool(grp["items"]))
            subs = [_norm_item(x) for x in grp["items"]]
            if len(subs) == 1 and not subs[0]["items"]:
                self.p(subs[0]["text"], align=L, left=HANG)
                continue
            for j, sub in enumerate(subs):
                self.p(f"{LETTERS[j]}. {sub['text']}", align=L, left=HANG, hang=HANG)
                for k, s3 in enumerate(sub["items"]):
                    s3 = _norm_item(s3)
                    self.p(f"{SMALL_ROMAN[k]}. {s3['text']}", align=L,
                           left=2 * HANG, hang=0.8)

    def s10_facilities(self):
        d = self.d
        self.heading("10. Facilities", hang=0.8)
        self.heading("a. Existing", left=HANG)
        self.p(d.get("facilities_existing", ""), align=J, left=HANG)
        self.heading("b. Additional", left=HANG)
        self.p(d.get("facilities_additional", ""), align=J, left=HANG)

    def s11_duration(self):
        self.heading("11. Duration of study", hang=0.8)
        self.p(self.d.get("duration", "Four semesters"), align=L, left=HANG)

    def s12_finance(self):
        """Borderless 2-column table so long item names wrap while the
        ': Rs.' column stays aligned."""
        self.heading("12. Financial estimate", hang=0.8)
        items = self.d.get("finance", [])
        rows, total = [], 0
        for i, it in enumerate(items):
            amt = _amount(it.get("amount", 0))
            total += amt
            rows.append((f"{LETTERS[i]}. {it.get('item','')}", f": Rs. {amt}"))
        rows.append(("Total", f": Rs. {total}"))
        widths = [COL_W - HANG - 2.4, 2.4]
        t = self.doc.add_table(rows=len(rows), cols=2)
        t.autofit = False
        tblPr = t._tbl.tblPr
        for tag, attrs in (("w:tblW", {"w:w": str(int(Cm(sum(widths)).twips)), "w:type": "dxa"}),
                           ("w:tblLayout", {"w:type": "fixed"}),
                           ("w:tblInd", {"w:w": str(int(Cm(HANG).twips)), "w:type": "dxa"})):
            el = OxmlElement(tag)
            for k, v in attrs.items():
                el.set(qn(k), v)
            tblPr.append(el)
        cm = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "0"); e.set(qn("w:type"), "dxa")
            cm.append(e)
        tblPr.append(cm)
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(Cm(w).twips)))
        for r, (label, amt) in enumerate(rows):
            for c, (txt, w) in enumerate(((label, widths[0]), (amt, widths[1]))):
                cell = t.cell(r, c)
                cell.width = Cm(w)
                par = cell.paragraphs[0]
                par.alignment = L
                if c == 0 and r < len(rows) - 1:
                    par.paragraph_format.left_indent = Cm(HANG)
                    par.paragraph_format.first_line_indent = Cm(-HANG)
                if c == 0 and r == len(rows) - 1:
                    par.paragraph_format.left_indent = Cm(HANG)
                add_rich(par, txt)

    def signatures(self):
        d = self.d
        self.blank()
        if self.sd:
            self.p("Sd/-", align=R)
        self.p("Signature of student", align=R, bold=True)
        self.blank()
        self.p("Project coordination group", align=L, bold=True, keep_next=True)
        self.p(d.get("pc_group", ""), align=L)
        self.blank()
        if self.sd:
            self.p("Sd/-", align=R)
        self.p("Signature of major advisor", align=R, bold=True)
        self.blank()
        self.p(f"Place: {d.get('place','')}", align=L)
        self.p(f"Date: {d.get('date','')}", align=L)

    def committee(self):
        self.p("Name, address and signature of other members of the advisory "
               "committee", align=L, bold=True, keep_next=True)
        for i, m in enumerate(self.d.get("committee", []), 1):
            par = self.p("", align=L, hang=HANG, keep_next=True)
            self.right_tab(par)
            add_rich(par, f"{i}. {m.get('name','')}" + ("\tSd/-" if self.sd else ""))
            addr = _lines(m.get("address", ""))
            for k, line in enumerate(addr):
                self.p(line, align=L, left=HANG, keep_next=(k < len(addr) - 1))

    def s13_references(self):
        self.heading("13. References", hang=0.8)
        style = self.d.get("ref_style", "KVASU thesis style")
        refs = sorted(self.d.get("references", []), key=sort_key)
        for ref in refs:
            par = self.p("", align=J, hang=1.0)
            for text, fmt in format_reference(ref, style):
                run = par.add_run(text)
                _font(run, bold=fmt.get("bold", False), italic=fmt.get("italic", False))

    def s14_timeframe(self):
        self.heading("14. Time frame of work", hang=0.8)
        for i, sem in enumerate(self.d.get("timeframe", [])):
            self.p(f"Semester {ROMAN[i]}", align=L, bold=True, left=HANG, keep_next=True)
            for j, task in enumerate(sem, 1):
                self.p(f"{j}. {task}", align=L, left=HANG, hang=HANG)

    def certificate(self):
        d = self.d
        self.p("CERTIFICATE", align=C, bold=True, keep_next=True)
        self.p("Certified that the research project has been formulated observing "
               "the stipulations laid down under the Prevention of Cruelty to "
               "Animals Act (Amendment, 1998).", align=J, first=1.0)
        self.blank()
        if self.sd:
            self.p("Sd/-", align=R)
        self.p(d.get("advisor_name", ""), align=R)
        self.p("Major advisor", align=R)
        self.blank()
        self.p(f"Place: {d.get('place','')}", align=L)
        self.p(f"Date: {d.get('date','')}", align=L)


# ---------------------------------------------------------------- utils
def _paras(text) -> list[str]:
    if isinstance(text, list):
        return [t.strip() for t in text if t and t.strip()]
    return [re.sub(r"\s*\n\s*", " ", p).strip()
            for p in re.split(r"\n\s*\n", text or "") if p.strip()]


def _lines(text) -> list[str]:
    if isinstance(text, list):
        return [t for t in text if t.strip()]
    return [l.strip() for l in (text or "").splitlines() if l.strip()]


def _amount(v) -> int:
    try:
        return int(round(float(str(v).replace(",", "").replace("Rs.", "").strip() or 0)))
    except ValueError:
        return int(re.sub(r"[^\d]", "", str(v)) or 0)


def _norm_item(x):
    if isinstance(x, str):
        return {"text": x, "items": []}
    return {"text": x.get("text", ""), "items": x.get("items", [])}


def build_document(data: dict, out_path: str) -> str:
    return Builder(data).build(out_path)
